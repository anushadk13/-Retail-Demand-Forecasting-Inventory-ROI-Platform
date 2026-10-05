"""Translate forecast accuracy into dollars and write the Excel business case.

Model (per category, all drivers visible in the workbook):
  safety stock      SS = z * sum_i(sigma_i) * sqrt(LT)       sigma_i = std of daily forecast error
  accuracy dividend D  = realisation * (1 - sigma_model / sigma_baseline)
  policy            a share `take_as_inventory` of D is taken as lower safety stock,
                    the rest is kept as extra service (fewer stockouts)
  stockouts         expected shortage per cycle = sigma_L * G(z),  G(z) = phi(z) - z(1 - Phi(z))
  savings           holding cost on released safety stock  +  margin recovered on avoided shortages

Baseline = 28-day moving average (what a planner without a model typically does), which is a
more conservative reference than seasonal-naive.
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.stats import norm

from src.config import ROOT, SCHEMA

BASELINE, MODEL = "moving_avg_28", "lightgbm"
DEFAULTS = dict(service_level=0.95, lead_time=7, holding_cost=0.25, gross_margin=0.30,
                take_as_inventory=0.5)
SCENARIOS = {"Conservative": 0.50, "Expected": 0.75, "Optimistic": 1.00}  # realisation of backtest gain
OUT_PATH = ROOT / "business_case" / "inventory_roi_model.xlsx"


def category_inputs(con) -> pd.DataFrame:
    """Per category: sum of per-series daily error sigma (baseline & model), avg price, annual units."""
    df = con.execute(f"""
        select f.product_key, f.store_key, p.category, f.model_name,
               f.forecast_units - f.actual_units as err, f.actual_units
        from {SCHEMA}.fact_forecast f join {SCHEMA}.dim_product p using (product_key)
        where f.model_name in ('{BASELINE}', '{MODEL}')""").df()
    sig = (df.groupby(["category", "product_key", "store_key", "model_name"])["err"].std()
             .rename("sigma").reset_index())
    agg = sig.groupby(["category", "model_name"])["sigma"].sum().unstack()
    px = con.execute(f"""
        select p.category, sum(f.revenue) / nullif(sum(f.units_sold), 0) as avg_price,
               sum(f.units_sold) * 365.0 / count(distinct f.date_key) as annual_units
        from {SCHEMA}.fact_sales f join {SCHEMA}.dim_product p using (product_key)
        group by 1""").df().set_index("category")
    out = pd.DataFrame({"sigma_sum_baseline": agg[BASELINE], "sigma_sum_model": agg[MODEL]}).join(px)
    return out.reset_index()


def G(z):  # standard normal loss function
    return norm.pdf(z) - z * (1 - norm.cdf(z))


def scenario_savings(inp: pd.DataFrame, realisation: float, p: dict | None = None,
                     improvement: float | None = None) -> pd.DataFrame:
    """Python mirror of the workbook formulas (used for tests and to cross-check the xlsx)."""
    p = {**DEFAULTS, **(p or {})}
    z = norm.ppf(p["service_level"])
    rows = []
    for r in inp.itertuples():
        a = r.sigma_sum_baseline * np.sqrt(p["lead_time"])
        d = improvement if improvement is not None else realisation * (1 - r.sigma_sum_model / r.sigma_sum_baseline)
        b = a * (1 - d)
        z_new = z * (1 - p["take_as_inventory"] * d) / (1 - d)
        unit_cost = r.avg_price * (1 - p["gross_margin"])
        ss_cut = z * a * p["take_as_inventory"] * d
        short_base, short_new = a * G(z), b * G(z_new)
        cycles = 365 / p["lead_time"]
        hold = ss_cut * unit_cost * p["holding_cost"]
        stockout = (short_base - short_new) * cycles * r.avg_price * p["gross_margin"]
        rows.append(dict(category=r.category, ss_base=z * a, ss_reduction_pct=p["take_as_inventory"] * d,
                         ss_cut_value=ss_cut * unit_cost, stockout_reduction_pct=1 - short_new / short_base,
                         holding_savings=hold, stockout_savings=stockout, total=hold + stockout,
                         ss_base_value=z * a * unit_cost, shortage_base_value=short_base * cycles * r.avg_price * p["gross_margin"]))
    return pd.DataFrame(rows)


def scenario_summary(inp: pd.DataFrame, p: dict | None = None) -> pd.DataFrame:
    out = []
    for name, real in SCENARIOS.items():
        s = scenario_savings(inp, real, p)
        out.append(dict(scenario=name,
                        ss_reduction_pct=s.ss_cut_value.sum() / s.ss_base_value.sum(),
                        stockout_reduction_pct=1 - (s.shortage_base_value.sum() - s.stockout_savings.sum()) / s.shortage_base_value.sum(),
                        annual_savings=s.total.sum()))
    return pd.DataFrame(out)


# ------------------------------------------------------------------ workbook
HDR = PatternFill("solid", fgColor="1F3A5F")
INP = PatternFill("solid", fgColor="FFF4CC")


def _hdr(ws, row, labels):
    for c, v in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def build_workbook(inp: pd.DataFrame, path=OUT_PATH) -> None:
    wb = Workbook()
    a = wb.active
    a.title = "Assumptions"
    a["A1"], a["A1"].font = "Inventory ROI model - assumptions (yellow cells are inputs)", Font(bold=True, size=13)
    labels = [("Service level", DEFAULTS["service_level"], "0%", "Target in-stock probability per replenishment cycle"),
              ("z (service level)", "=NORMSINV(B3)", "0.000", "Derived"),
              ("Lead time (days)", DEFAULTS["lead_time"], "0", "Replenishment lead time = review period"),
              ("Holding cost (% of inventory value / yr)", DEFAULTS["holding_cost"], "0%", "Capital + storage + shrink"),
              ("Gross margin on lost sales", DEFAULTS["gross_margin"], "0%", "Margin lost on each unit short; unit cost = price x (1 - margin)"),
              ("Share of accuracy gain taken as lower inventory", DEFAULTS["take_as_inventory"], "0%", "Remainder is kept as extra service (fewer stockouts)"),
              ("Replenishment cycles per year", "=365/B5", "0.0", "Derived")]
    for i, (lab, val, fmt, note) in enumerate(labels, start=3):
        a.cell(row=i, column=1, value=lab)
        c = a.cell(row=i, column=2, value=val)
        c.number_format = fmt
        if not str(val).startswith("="):
            c.fill = INP
        a.cell(row=i, column=3, value=note)
    a["A11"], a["A11"].font = "Safety stock = z x sum(sigma_i) x sqrt(lead time); sigma_i = std dev of daily forecast error of series i.", Font(italic=True)
    a["A12"] = f"Baseline = {BASELINE}; model = {MODEL}. Error sigmas come from the 13-fold rolling backtest (src/evaluation/roi.py)."
    a["A13"] = "Realisation: share of the backtest accuracy gain assumed to hold in production (Conservative 50% / Expected 75% / Optimistic 100%)."
    a["A14"] = "NOTE: dollar values are computed on SYNTHETIC sample data unless the real M5 dataset was loaded."
    a.column_dimensions["A"].width, a.column_dimensions["B"].width, a.column_dimensions["C"].width = 48, 14, 70

    d = wb.create_sheet("Data")
    _hdr(d, 1, ["Category", "Sum of series sigma - baseline", "Sum of series sigma - model",
                "Avg selling price ($/unit)", "Annual units"])
    for i, r in enumerate(inp.itertuples(), start=2):
        for c, v in enumerate([r.category, r.sigma_sum_baseline, r.sigma_sum_model, r.avg_price, r.annual_units], 1):
            d.cell(row=i, column=c, value=v)
    n = len(inp)
    for col, w in zip("ABCDE", (14, 22, 22, 18, 14)):
        d.column_dimensions[col].width = w

    s = wb.create_sheet("Scenarios")
    _hdr(s, 1, ["Scenario", "Realisation of backtest gain", "Category", "Accuracy dividend D", "Safety stock before ($)",
                "Safety stock reduction (%)", "Safety stock released ($)", "Effective z after", "Stockout reduction (%)",
                "Holding-cost savings ($/yr)", "Stockout margin recovered ($/yr)", "Total savings ($/yr)",
                "Shortage cost before ($/yr)"])
    row = 2
    summary_rows = {}
    for name, real in SCENARIOS.items():
        first = row
        for k in range(n):
            dr = k + 2  # Data row
            A = f"(Data!$B{dr}*SQRT(Assumptions!$B$5))"
            cost = f"(Data!$D{dr}*(1-Assumptions!$B$7))"
            s.cell(row=row, column=1, value=name)
            s.cell(row=row, column=2, value=real).number_format = "0%"
            s.cell(row=row, column=3, value=f"=Data!A{dr}")
            s.cell(row=row, column=4, value=f"=B{row}*(1-Data!C{dr}/Data!B{dr})").number_format = "0.0%"
            s.cell(row=row, column=5, value=f"=Assumptions!$B$4*{A}*{cost}").number_format = "#,##0"
            s.cell(row=row, column=6, value=f"=Assumptions!$B$8*D{row}").number_format = "0.0%"
            s.cell(row=row, column=7, value=f"=E{row}*F{row}").number_format = "#,##0"
            s.cell(row=row, column=8, value=f"=Assumptions!$B$4*(1-Assumptions!$B$8*D{row})/(1-D{row})").number_format = "0.000"
            G0 = "(NORMDIST(Assumptions!$B$4,0,1,FALSE)-Assumptions!$B$4*(1-NORMSDIST(Assumptions!$B$4)))"
            G1 = f"(NORMDIST(H{row},0,1,FALSE)-H{row}*(1-NORMSDIST(H{row})))"
            s.cell(row=row, column=9, value=f"=1-((1-D{row})*{G1})/{G0}").number_format = "0.0%"
            s.cell(row=row, column=10, value=f"=G{row}*Assumptions!$B$6").number_format = "#,##0"
            margin = f"(Data!$D{dr}*Assumptions!$B$7)"
            s.cell(row=row, column=13, value=f"={A}*{G0}*Assumptions!$B$9*{margin}").number_format = "#,##0"
            s.cell(row=row, column=11, value=f"=M{row}*I{row}").number_format = "#,##0"
            s.cell(row=row, column=12, value=f"=J{row}+K{row}").number_format = "#,##0"
            row += 1
        summary_rows[name] = (first, row - 1)
    for col in "ABCDEFGHIJKLM":
        s.column_dimensions[col].width = 18
    s.row_dimensions[1].height = 45

    t = wb.create_sheet("Summary", 0)
    t["A1"], t["A1"].font = "Annual savings from better forecasting", Font(bold=True, size=13)
    _hdr(t, 3, ["Scenario", "Safety stock reduction", "Stockout reduction", "Annual savings ($)"])
    for i, (name, (f, l)) in enumerate(summary_rows.items(), start=4):
        t.cell(row=i, column=1, value=name)
        t.cell(row=i, column=2, value=f"=SUM(Scenarios!G{f}:G{l})/SUM(Scenarios!E{f}:E{l})").number_format = "0.0%"
        t.cell(row=i, column=3, value=f"=SUM(Scenarios!K{f}:K{l})/SUM(Scenarios!M{f}:M{l})").number_format = "0.0%"
        t.cell(row=i, column=4, value=f"=SUM(Scenarios!L{f}:L{l})").number_format = "$#,##0"
    t["A9"] = "Sensitivity: annual savings ($) by accuracy improvement (sigma reduction, rows) x holding cost (columns)"
    t["A9"].font = Font(bold=True)
    hold_grid = [0.10, 0.15, 0.20, 0.25, 0.30, 0.40]
    imp_grid = [0.05, 0.10, 0.15, 0.20, 0.25, 0.30]
    t.cell(row=10, column=1, value="Improvement \\ Holding cost").font = Font(bold=True)
    for j, h in enumerate(hold_grid, start=2):
        c = t.cell(row=10, column=j, value=h)
        c.number_format, c.font = "0%", Font(bold=True)
    for i, imp in enumerate(imp_grid, start=11):
        c = t.cell(row=i, column=1, value=imp)
        c.number_format, c.font = "0%", Font(bold=True)
        for j in range(2, 2 + len(hold_grid)):
            col = t.cell(row=10, column=j).column_letter
            terms = []
            for k in range(n):
                dr = k + 2
                A = f"(Data!$B{dr}*SQRT(Assumptions!$B$5))"
                cost = f"(Data!$D{dr}*(1-Assumptions!$B$7))"
                margin = f"(Data!$D{dr}*Assumptions!$B$7)"
                zn = f"(Assumptions!$B$4*(1-Assumptions!$B$8*$A{i})/(1-$A{i}))"
                G0 = "(NORMDIST(Assumptions!$B$4,0,1,FALSE)-Assumptions!$B$4*(1-NORMSDIST(Assumptions!$B$4)))"
                G1 = f"(NORMDIST({zn},0,1,FALSE)-{zn}*(1-NORMSDIST({zn})))"
                hold = f"Assumptions!$B$4*{A}*Assumptions!$B$8*$A{i}*{cost}*{col}$10"
                short = f"{A}*Assumptions!$B$9*{margin}*({G0}-(1-$A{i})*{G1})"
                terms.append(f"{hold}+{short}")
            t.cell(row=i, column=j, value="=" + "+".join(terms)).number_format = "$#,##0"
    t.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        t.column_dimensions[col].width = 18
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    from src.config import connect
    inp = category_inputs(connect(read_only=True))
    build_workbook(inp)
    pd.options.display.float_format = "{:,.3f}".format
    print(inp)
    print(scenario_summary(inp))
    print(f"wrote {OUT_PATH}")
